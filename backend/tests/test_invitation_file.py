"""Invitation files: one shape each, and no extra columns."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.invitation_file import InvitationFileError, parse_invitation_file

EXAMPLES = Path(__file__).resolve().parents[2] / "docs" / "tool"


def test_json_example_keeps_only_emails():
    parsed = parse_invitation_file("invitations.json", (EXAMPLES / "invitations.json").read_bytes())
    assert parsed.format == "json"
    assert parsed.emails == ["pessoa.um@example.com", "pessoa.dois@example.com"]


def test_json_rejects_a_bare_array_and_any_other_shape():
    with pytest.raises(InvitationFileError) as bare:
        parse_invitation_file("list.json", b'[{"email":"a@example.com"}]')
    assert bare.value.code == "bad_file"

    with pytest.raises(InvitationFileError) as extra:
        parse_invitation_file(
            "list.json",
            b'{"invitations":[{"email":"a@example.com"}],"company":"x"}',
        )
    assert extra.value.code == "unknown_columns"
    assert extra.value.columns == ["company"]


def test_json_rejects_a_name_and_stores_nothing_from_that_file():
    with pytest.raises(InvitationFileError) as error:
        parse_invitation_file(
            "list.json",
            b'{"invitations":[{"email":"a@example.com","name":"Ada"}]}',
        )
    assert error.value.code == "unknown_columns"
    assert error.value.columns == ["name"]


def test_csv_example_and_semicolon_header():
    comma = parse_invitation_file("invitations.csv", (EXAMPLES / "invitations.csv").read_bytes())
    assert comma.format == "csv"
    assert comma.emails == ["pessoa.um@example.com", "pessoa.dois@example.com"]

    semi = parse_invitation_file(
        "invitations.semicolon.csv",
        (EXAMPLES / "invitations.semicolon.csv").read_bytes(),
    )
    assert semi.emails == ["pessoa.tres@example.com"]


def test_csv_rejects_unknown_columns_and_ignores_empty_rows():
    with pytest.raises(InvitationFileError) as named:
        parse_invitation_file("people.csv", b"email,name\nada@example.com,Ada\n")
    assert named.value.code == "unknown_columns"
    assert named.value.columns == ["name"]

    with pytest.raises(InvitationFileError) as nome:
        parse_invitation_file("people.csv", "nome\nada@example.com\n".encode())
    assert nome.value.code == "unknown_columns"

    parsed = parse_invitation_file(
        "people.csv",
        "email\n\npessoa.um@example.com\n,\nPESSOA.UM@example.com\nnot-an-email\n".encode(),
    )
    assert parsed.emails == ["pessoa.um@example.com"]
    assert parsed.duplicates_in_file == ["pessoa.um@example.com"]
    assert parsed.invalid == ["not-an-email"]
    assert parsed.ignored_empty == 2


def test_csv_rejects_bytes_that_are_not_utf8():
    with pytest.raises(InvitationFileError) as error:
        parse_invitation_file("people.csv", b"email\n\xff\n")
    assert error.value.code == "bad_file"


def test_xlsx_example_matches_the_csv_example():
    parsed = parse_invitation_file("invitations.xlsx", (EXAMPLES / "invitations.xlsx").read_bytes())
    assert parsed.format == "xlsx"
    assert parsed.emails == ["pessoa.um@example.com", "pessoa.dois@example.com"]


def test_xlsx_reads_the_first_sheet_only_and_rejects_a_second_column():
    workbook = Workbook()
    first = workbook.active
    first.title = "convites"
    first.append(["email"])
    first.append(["pessoa.um@example.com"])
    first.append([""])
    second = workbook.create_sheet("nomes")
    second.append(["nome"])
    second.append(["Ada"])
    data = _workbook_bytes(workbook)

    parsed = parse_invitation_file("lista.xlsx", data)
    assert parsed.format == "xlsx"
    assert parsed.emails == ["pessoa.um@example.com"]
    assert parsed.ignored_empty == 1

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["email", "cargo"])
    sheet.append(["pessoa.um@example.com", "analista"])
    with pytest.raises(InvitationFileError) as error:
        parse_invitation_file("lista.xlsx", _workbook_bytes(workbook))
    assert error.value.code == "unknown_columns"
    assert "cargo" in error.value.columns


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
