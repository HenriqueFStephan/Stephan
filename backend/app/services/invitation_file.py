"""Parse one invitation list. The only accepted field is email.

JSON has one shape. CSV is UTF-8 with the header `email`. Excel is the first
worksheet of an .xlsx file, same header on the first row. Unknown columns
reject the whole file so a name is never stored by accident.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from io import BytesIO

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook

MAX_BYTES = 1_000_000
MAX_ROWS = 5_000


class InvitationFileError(Exception):
    def __init__(self, code: str, columns: list[str] | None = None) -> None:
        self.code = code
        self.columns = columns or []
        super().__init__(code)


@dataclass
class ParsedInvitations:
    format: str
    emails: list[str] = field(default_factory=list)
    ignored_empty: int = 0
    duplicates_in_file: list[str] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)


def parse_invitation_file(filename: str, data: bytes) -> ParsedInvitations:
    if len(data) > MAX_BYTES:
        raise InvitationFileError("bad_file")
    kind = _kind(filename, data)
    if kind == "xlsx":
        rows = _xlsx_rows(data)
        parsed = _from_rows(rows)
        parsed.format = "xlsx"
        return parsed
    if kind == "json":
        parsed = _from_json(data)
        parsed.format = "json"
        return parsed
    parsed = _from_csv(data)
    parsed.format = "csv"
    return parsed


def _kind(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if data[:2] == b"PK" or name.endswith(".xlsx"):
        return "xlsx"
    stripped = data.lstrip(b"\xef\xbb\xbf \t\r\n")
    if stripped.startswith(b"{") or name.endswith(".json"):
        return "json"
    return "csv"


def _from_json(data: bytes) -> ParsedInvitations:
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise InvitationFileError("bad_file") from exc
    if not isinstance(payload, dict):
        raise InvitationFileError("bad_file")
    extra_root = sorted(set(payload) - {"invitations"})
    if extra_root:
        raise InvitationFileError("unknown_columns", extra_root)
    if set(payload) != {"invitations"} or not isinstance(payload["invitations"], list):
        raise InvitationFileError("bad_file")
    rows: list[list[str]] = [["email"]]
    for item in payload["invitations"]:
        if not isinstance(item, dict):
            raise InvitationFileError("bad_file")
        extra = sorted(str(key) for key in item if key != "email")
        if extra:
            raise InvitationFileError("unknown_columns", extra)
        if "email" not in item:
            raise InvitationFileError("bad_file")
        value = item["email"]
        if value is None:
            rows.append([""])
        elif isinstance(value, str):
            rows.append([value])
        else:
            rows.append([str(value)])
    return _from_rows(rows)


def _from_csv(data: bytes) -> ParsedInvitations:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvitationFileError("bad_file") from exc
    lines = text.splitlines()
    if not lines or not lines[0].strip():
        raise InvitationFileError("bad_file")
    delimiter = _delimiter(lines[0])
    rows: list[list[str]] = []
    for line in lines:
        if not line.strip():
            rows.append([""])
            continue
        try:
            cells = next(csv.reader([line], delimiter=delimiter))
        except csv.Error as exc:
            raise InvitationFileError("bad_file") from exc
        rows.append(cells)
    return _from_rows(rows)


def _delimiter(header_line: str) -> str:
    if ";" in header_line:
        _require_email_header(_header_cells(header_line, ";"))
        return ";"
    _require_email_header(_header_cells(header_line, ","))
    return ","


def _require_email_header(cells: list[str]) -> None:
    if cells == ["email"]:
        return
    unknown = [cell for cell in cells if cell and cell != "email"]
    if unknown:
        raise InvitationFileError("unknown_columns", unknown)
    raise InvitationFileError("bad_file")


def _header_cells(line: str, delimiter: str) -> list[str]:
    cells = next(csv.reader([line], delimiter=delimiter))
    while cells and cells[-1].strip() == "":
        cells.pop()
    return [cell.strip() for cell in cells]


def _xlsx_rows(data: bytes) -> list[list[str]]:
    if data[:2] != b"PK":
        raise InvitationFileError("bad_file")
    try:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:
        raise InvitationFileError("bad_file") from exc
    try:
        if not workbook.worksheets:
            raise InvitationFileError("bad_file")
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        return rows
    finally:
        workbook.close()


def _from_rows(rows: list[list[str]]) -> ParsedInvitations:
    if not rows:
        raise InvitationFileError("bad_file")
    _require_email_header(_trim_row(rows[0]))
    if len(rows) - 1 > MAX_ROWS:
        raise InvitationFileError("bad_file")

    parsed = ParsedInvitations(format="csv")
    seen: set[str] = set()
    for raw in rows[1:]:
        _reject_side_values(raw)
        kind, value = _classify_email(raw[0] if raw else "")
        if kind == "empty":
            parsed.ignored_empty += 1
        elif kind == "invalid":
            parsed.invalid.append(value)
        elif value in seen:
            parsed.duplicates_in_file.append(value)
        else:
            seen.add(value)
            parsed.emails.append(value)
    return parsed


def _trim_row(row: list[str]) -> list[str]:
    cells = list(row)
    while cells and cells[-1].strip() == "":
        cells.pop()
    return [cell.strip() for cell in cells]


def _reject_side_values(row: list[str]) -> None:
    extras = [cell.strip() for cell in row[1:] if cell.strip()]
    if extras:
        raise InvitationFileError("unknown_columns", ["column_2"])


def _classify_email(value: str) -> tuple[str, str]:
    text = value.strip()
    if not text:
        return "empty", ""
    if any(ch in text for ch in "\r\n\t"):
        return "invalid", text[:200]
    try:
        checked = validate_email(text, check_deliverability=False)
    except EmailNotValidError:
        return "invalid", text[:200]
    return "ok", checked.normalized.casefold()
